# -*- coding: utf-8 -*-
"""xlsx 构造/解析辅助（EasyExcel 写出的表头是中文列名）。"""
import io

from openpyxl import Workbook, load_workbook

GATE_HEADERS = ["人员编号", "部门", "姓名", "手机号", "身份证号"]


def build_gate_person_xlsx(rows):
    """rows: [(code, dept, name, phone, idCard), ...]；dept 传 None 表示留空。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(GATE_HEADERS)
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_sheet(content):
    """返回 (sheetname, [row_values...])，含表头行。"""
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb[wb.sheetnames[0]]
    return wb.sheetnames[0], [[c for c in row] for row in ws.iter_rows(values_only=True)]


def data_rows(content):
    """去掉表头后的数据行。"""
    _, rows = read_sheet(content)
    return [r for r in rows[1:] if any(v not in (None, "") for v in r)]
